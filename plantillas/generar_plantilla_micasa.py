#!/usr/bin/env python3
"""Genera la planilla de NOVEDADES MENSUALES que el estudio envía al cliente
para que complete los datos necesarios para liquidar los sueldos.

Empresa: MI CASA SOCIEDAD ANONIMA. Contempla todos los conceptos que se
liquidan (novedades variables del mes + altas de personal).

Uso: python3 generar_plantilla_micasa.py [salida.xlsx]
"""

import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

EMPRESA = "MI CASA SOCIEDAD ANONIMA"

# Paleta AsysTax
NAVY = "0B1B3A"
PRIMARY = "1E5BFF"
BLUE_SOFT = "E7EEFF"
GREEN_SOFT = "E7F6EC"
RED_SOFT = "FBEAEA"
GREY_SOFT = "F2F4F8"
WHITE = "FFFFFF"

thin = Side(style="thin", color="C7CEDB")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def title_cell(ws, cell, text, size=14, color=NAVY, bold=True):
    ws[cell] = text
    ws[cell].font = Font(name="Calibri", size=size, bold=bold, color=color)


def header_row(ws, row, headers, fill=PRIMARY, font_color=WHITE, start_col=1):
    for i, (label, width, group) in enumerate(headers):
        col = start_col + i
        c = ws.cell(row=row, column=col, value=label)
        c.font = Font(bold=True, color=font_color, size=9)
        c.fill = PatternFill("solid", fgColor=group or fill)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
        ws.column_dimensions[get_column_letter(col)].width = width


def build():
    wb = Workbook()

    # ───────────────────────── INSTRUCTIVO ─────────────────────────
    ins = wb.active
    ins.title = "Instructivo"
    ins.sheet_view.showGridLines = False
    ins.column_dimensions["A"].width = 3
    ins.column_dimensions["B"].width = 30
    ins.column_dimensions["C"].width = 85

    title_cell(ins, "B2", "AsysTax. Sueldos", size=18, color=PRIMARY)
    title_cell(ins, "B3", "Planilla de novedades mensuales", size=13, color=NAVY)
    title_cell(ins, "B4", f"Empresa: {EMPRESA}", size=11, color=NAVY)
    ins["B5"] = "GRO Consultores & Asociados · jortiz@gro.com.uy"
    ins["B5"].font = Font(size=9, italic=True, color="6B7280")

    title_cell(ins, "B7", "Cómo completar", size=12, color=PRIMARY)
    instrucciones = [
        ("Mes a liquidar", "Indicá el mes y año arriba en la hoja 'Novedades del mes'."),
        ("Una fila por empleado", "En 'Novedades del mes' cargá una fila por cada persona activa."),
        ("Días trabajados", "Normalmente 30. Si ingresó/egresó a mitad de mes, poné los días reales."),
        ("Faltas", "Cantidad de DÍAS de falta. El monto se calcula solo (jornal = básico/30)."),
        ("Horas tarde", "Cantidad de HORAS de llegada tarde (hora = jornal/8)."),
        ("Licencia gozada", "Días de licencia que se toma EN EL MES (se paga en la mensualidad)."),
        ("Salario vacacional (días)", "Días de licencia a pagar como salario vacacional (liquidación aparte, exento)."),
        ("Horas extra", "Diurnas y nocturnas por separado (cantidad de horas)."),
        ("Descansos trabajados", "Cantidad de descansos trabajados (se paga un jornal por cada uno)."),
        ("Comisiones", "Monto en $ de comisiones del mes (gravado)."),
        ("Viáticos gravados", "Monto en $ de viáticos GRAVADOS (integran aportes e IRPF)."),
        ("Viáticos no gravados", "Monto en $ de viáticos NO gravados (suman al líquido sin descuentos)."),
        ("Retención judicial %", "Porcentaje de embargo sobre el total de haberes (si corresponde)."),
        ("Adelantos", "Monto en $ adelantado en el mes (se descuenta del líquido)."),
        ("Otros haberes / descuentos", "Detalle + monto de cualquier otro concepto no listado."),
        ("Observaciones", "Cualquier aclaración para el estudio."),
    ]
    r = 8
    for campo, desc in instrucciones:
        ins.cell(row=r, column=2, value=campo).font = Font(bold=True, size=10, color=NAVY)
        ins.cell(row=r, column=3, value=desc).font = Font(size=10, color="374151")
        ins.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1

    r += 1
    title_cell(ins, f"B{r}", "Automático (no completar)", size=12, color=PRIMARY)
    r += 1
    autos = [
        ("Sueldo básico", "Lo tiene registrado el estudio en el contrato de cada empleado."),
        ("Prima por antigüedad", f"En {EMPRESA} es 10% del sueldo básico, automática. Solo avisá si cambia."),
        ("Aportes y IRPF", "BPS, FONASA, FRL e IRPF se calculan solos según la normativa vigente."),
        ("Altas de personal", "Si ingresa alguien nuevo, completá la hoja 'Altas (nuevos empleados)'."),
    ]
    for campo, desc in autos:
        ins.cell(row=r, column=2, value=campo).font = Font(bold=True, size=10, color=NAVY)
        ins.cell(row=r, column=3, value=desc).font = Font(size=10, color="374151")
        ins.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1

    # ───────────────────── NOVEDADES DEL MES ─────────────────────
    nov = wb.create_sheet("Novedades del mes")
    nov.sheet_view.showGridLines = False

    title_cell(nov, "A1", f"Novedades mensuales — {EMPRESA}", size=13, color=NAVY)
    nov["A2"] = "Mes a liquidar:"
    nov["A2"].font = Font(bold=True, color=NAVY, size=10)
    nov["B2"] = "  (ej. Agosto 2026)"
    nov["B2"].font = Font(italic=True, color="9CA3AF", size=10)
    nov["B2"].fill = PatternFill("solid", fgColor="FFF7CC")
    nov["B2"].border = border

    # (label, width, group-color)
    cols = [
        ("Apellido y Nombre", 26, NAVY),
        ("C.I.", 12, NAVY),
        ("Días trabajados", 10, PRIMARY),
        ("Faltas (días)", 9, RED_SOFT),
        ("Horas tarde", 9, RED_SOFT),
        ("Licencia gozada (días)", 11, PRIMARY),
        ("Salario vacacional (días)", 11, PRIMARY),
        ("H. extra diurnas", 9, GREEN_SOFT),
        ("H. extra nocturnas", 9, GREEN_SOFT),
        ("Descansos trabajados", 10, GREEN_SOFT),
        ("Comisiones $", 11, GREEN_SOFT),
        ("Viáticos gravados $", 12, GREEN_SOFT),
        ("Viáticos NO gravados $", 12, GREEN_SOFT),
        ("Retención judicial %", 10, RED_SOFT),
        ("Adelantos $", 11, RED_SOFT),
        ("Otros haberes (detalle)", 22, GREY_SOFT),
        ("Otros haberes $", 11, GREEN_SOFT),
        ("Otros descuentos (detalle)", 22, GREY_SOFT),
        ("Otros descuentos $", 11, RED_SOFT),
        ("Observaciones", 28, GREY_SOFT),
    ]
    # Grupos con color de encabezado (tinta suave -> texto oscuro salvo NAVY/PRIMARY)
    hrow = 4
    for i, (label, width, group) in enumerate(cols, start=1):
        c = nov.cell(row=hrow, column=i, value=label)
        dark = group in (NAVY, PRIMARY)
        c.font = Font(bold=True, size=9, color=WHITE if dark else NAVY)
        c.fill = PatternFill("solid", fgColor=group)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
        nov.column_dimensions[get_column_letter(i)].width = width
    nov.row_dimensions[hrow].height = 42

    # Filas en blanco para completar (con bordes y default de días = 30)
    FILAS = 40
    for row in range(hrow + 1, hrow + 1 + FILAS):
        for col in range(1, len(cols) + 1):
            cell = nov.cell(row=row, column=col)
            cell.border = border
            cell.font = Font(size=10)
            if col >= 3:
                cell.alignment = Alignment(horizontal="center")
        nov.cell(row=row, column=3, value=30)  # días trabajados por defecto

    nov.freeze_panes = "C5"
    nov.auto_filter.ref = f"A{hrow}:{get_column_letter(len(cols))}{hrow + FILAS}"

    # ───────────────────── ALTAS (NUEVOS) ─────────────────────
    alt = wb.create_sheet("Altas (nuevos empleados)")
    alt.sheet_view.showGridLines = False
    title_cell(alt, "A1", "Altas de personal — datos para el alta BPS y el contrato", size=12, color=NAVY)
    alt["A2"] = "Completar solo si ingresa una persona nueva."
    alt["A2"].font = Font(italic=True, color="6B7280", size=10)

    acols = [
        ("Apellido", 18), ("Nombre", 18), ("C.I.", 13),
        ("Fecha nacimiento", 14), ("Fecha ingreso", 14),
        ("Cargo", 18), ("Sector", 16),
        ("Sueldo nominal mensual $", 15), ("Tipo (Mensual/Jornalero)", 15),
        ("Jornal $ (si jornalero)", 13),
        ("Hijos a cargo (cant.)", 11), ("Cónyuge a cargo (Sí/No)", 12),
        ("Banco", 18), ("N° de cuenta", 20),
        ("Domicilio", 26), ("Teléfono", 14), ("Email", 22),
        ("Observaciones", 26),
    ]
    ahrow = 4
    for i, (label, width) in enumerate(acols, start=1):
        c = alt.cell(row=ahrow, column=i, value=label)
        c.font = Font(bold=True, size=9, color=WHITE)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
        alt.column_dimensions[get_column_letter(i)].width = width
    alt.row_dimensions[ahrow].height = 40

    AFILAS = 20
    for row in range(ahrow + 1, ahrow + 1 + AFILAS):
        for col in range(1, len(acols) + 1):
            cell = alt.cell(row=row, column=col)
            cell.border = border
            cell.font = Font(size=10)

    # Validaciones: Tipo y Cónyuge
    dv_tipo = DataValidation(type="list", formula1='"Mensual,Jornalero"', allow_blank=True)
    dv_si_no = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
    alt.add_data_validation(dv_tipo)
    alt.add_data_validation(dv_si_no)
    dv_tipo.add(f"I{ahrow+1}:I{ahrow+AFILAS}")   # Tipo
    dv_si_no.add(f"L{ahrow+1}:L{ahrow+AFILAS}")   # Cónyuge a cargo
    alt.freeze_panes = "A5"

    out = sys.argv[1] if len(sys.argv) > 1 else f"Plantilla Novedades - {EMPRESA}.xlsx"
    wb.save(out)
    print(f"Planilla generada: {out}")
    return out


if __name__ == "__main__":
    build()
